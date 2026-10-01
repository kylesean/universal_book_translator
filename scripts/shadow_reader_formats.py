#!/usr/bin/env python
"""Phase-1 acceptance: the native HTML/EPUB/DOCX readers (ADR-0001 §12 Q5).

Each reader must give every element a real ``Span.chars`` that slices exactly its
own text out of ``CanonicalSource.text``, and the ``Document`` must round-trip
through the bridge losslessly. The readers are exercised over fixtures built here
(a minimal DOCX via python-docx, an EPUB via zipfile, an HTML string), so the
harness needs no external sample files.

Exit 0 iff all three readers pass span exactness, round-trip and structure.
"""

from __future__ import annotations

import argparse
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from ubt.analyze.assemble import element_text
from ubt.analyze.bridge import blocks_from_document, document_from_blocks
from ubt.analyze.reader_docx import read_docx
from ubt.analyze.reader_epub import read_epub
from ubt.analyze.reader_html import read_html
from ubt.model.ast import Document

#: A 1x1 transparent PNG, for the DOCX picture fixture.
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360606060000000050001a5f645400000000049454e44ae426082"
)

_HTML = """<!doctype html><html><head><title>t</title></head><body>
<h1>Chapter One</h1>
<p>Hello <b>world</b>, this is prose.</p>
<ul><li>First item</li><li>Second item</li></ul>
<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>
<figure><img src="pic.png"><figcaption>Figure 1</figcaption></figure>
<pre>print("hi")</pre>
</body></html>"""

_EPUB_CONTAINER = """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf"
    media-type="application/oebps-package+xml"/></rootfiles>
</container>"""

_EPUB_OPF = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>T</dc:title></metadata>
  <manifest>
    <item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>
    <item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>"""

_EPUB_CH1 = """<html><body><h1>One</h1><p>First chapter.</p>
<img src="images/fig.png" alt="fig"/></body></html>"""
_EPUB_CH2 = """<html><body><h1>Two</h1><p>Second chapter.</p>
<ul><li>a</li><li>b</li></ul></body></html>"""


def _span_problems(document: Document) -> list[str]:
    problems: list[str] = []
    for element in document.elements:
        chars = element.span.chars
        if chars is None:
            problems.append(f"{element.id}: missing Span.chars")
            continue
        expected = element_text(element)
        got = document.source.text[chars[0] : chars[1]]
        if got != expected:
            problems.append(f"{element.id}: span slice {got[:30]!r} != {expected[:30]!r}")
    return problems


def _check(document: Document, expect: dict[str, int]) -> list[str]:
    problems = _span_problems(document)
    round_tripped = document_from_blocks(
        blocks_from_document(document), doc_id=document.source.doc_id
    )
    if len(round_tripped.elements) != len(document.elements):
        problems.append(
            f"round-trip element count {len(round_tripped.elements)} != {len(document.elements)}"
        )
    counts = Counter(element.kind.value for element in document.elements)
    for kind, minimum in expect.items():
        if counts.get(kind, 0) < minimum:
            problems.append(f"expected >= {minimum} {kind}, got {counts.get(kind, 0)}")
    return problems


def _write_docx(path: Path) -> None:
    from docx import Document as DocxDocument

    image = path.with_suffix(".png")
    image.write_bytes(_PNG)
    doc = DocxDocument()
    doc.add_heading("Chapter One", level=1)
    doc.add_paragraph("Hello world, this is prose.")
    doc.add_paragraph("First item", style="List Bullet")
    doc.add_paragraph("Second item", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.add_picture(str(image))
    doc.save(str(path))


def _write_epub(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr("META-INF/container.xml", _EPUB_CONTAINER)
        zf.writestr("OEBPS/content.opf", _EPUB_OPF)
        zf.writestr("OEBPS/ch1.xhtml", _EPUB_CH1)
        zf.writestr("OEBPS/ch2.xhtml", _EPUB_CH2)
        zf.writestr("OEBPS/images/fig.png", _PNG)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()

    cases: list[tuple[str, Document, dict[str, int]]] = []
    with tempfile.TemporaryDirectory(prefix="ubt-readers-") as tmp_str:
        tmp = Path(tmp_str)
        html_path = tmp / "sample.html"
        html_path.write_text(_HTML, encoding="utf-8")
        cases.append(
            (
                "html",
                read_html(html_path),
                {
                    "heading": 1,
                    "paragraph": 1,
                    "list_item": 2,
                    "table": 1,
                    "figure": 1,
                    "caption": 1,
                    "code_block": 1,
                },
            )
        )

        epub_path = tmp / "sample.epub"
        _write_epub(epub_path)
        cases.append(
            (
                "epub",
                read_epub(epub_path),
                {"heading": 2, "paragraph": 2, "figure": 1, "list_item": 2},
            )
        )

        docx_path = tmp / "sample.docx"
        _write_docx(docx_path)
        cases.append(
            (
                "docx",
                read_docx(docx_path),
                {"heading": 1, "paragraph": 1, "list_item": 2, "table": 1, "figure": 1},
            )
        )

        print("\nPhase-1 native readers (html / epub / docx)")
        failed = 0
        for name, document, expect in cases:
            problems = _check(document, expect)
            if problems:
                failed += 1
            status = "pass" if not problems else "FAIL"
            print(f"  {status:<5} elements={len(document.elements):<4} {name}")
            for problem in problems[:5]:
                print(f"        {problem}")

    print(f"\n  readers={len(cases)} failed={failed} -> {'PASS' if not failed else 'FAIL'}")
    return 0 if not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
