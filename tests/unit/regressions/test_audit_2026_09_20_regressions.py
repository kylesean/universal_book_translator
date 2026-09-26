"""Regressions from the 2026-09-20 deep audit (12-track review + adversarial verify).

Every guard reproduces its defect against real code — real ledger SQLite, real
DOCX XML, real DOM mining — rather than mocking the component under test. Three
of these defects stayed invisible precisely because their tests modelled the
component instead of using it.
"""

from __future__ import annotations

import base64
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast

import pytest

from tests.stage_ctx_factory import build_stage_ctx
from ubt.adapters.epub.adapter import BLOCK_TAGS, is_leaf_block
from ubt.adapters.html.adapter import HTMLAdapter
from ubt.adapters.markdown.adapter import MarkdownAdapter
from ubt.core.config import UBTConfig
from ubt.core.engine.ledger import SQLiteJobLedger
from ubt.core.engine.stages.export import run_export_stage
from ubt.core.exceptions import IntegrityViolationError
from ubt.core.ir.models import (
    BlockStatus,
    BlockType,
    BookManifest,
    DocumentIR,
    FlowID,
    IRBlock,
)
from ubt.core.language_profile import resolve_font_config
from ubt.core.presets import PRESETS, Preset, resolve_engine_params
from ubt.core.qe.comet_runner import QE_SCORE_FABRICATED, HeuristicQERunner
from ubt.core.qe.defect_taxonomy import (
    CRITICAL_DEFECT_MARKERS,
    ECHO_MARKER,
    NEAR_ECHO_MARKER,
    STRUCTURAL_DEFECT_MARKERS,
    has_structural_defect,
)
from ubt.core.qe.fast_pass import FastPassFilter
from ubt.core.validators.html_delta import HTMLDeltaValidator

_SOURCE_PARAGRAPH = (
    "The device operates in inversion when the gate exceeds the threshold "
    "voltage across the oxide layer here."
)


# --- EPUB / HTML leaf mining -------------------------------------------------


def _html_leaves(markup: str) -> list[str]:
    from bs4 import BeautifulSoup

    adapter = HTMLAdapter()
    soup = BeautifulSoup(markup, "html.parser")
    return [f"{tag.name}:{tag.get_text(' ', strip=True)}" for tag in adapter._leaf_blocks(soup)]


def test_wrapper_div_is_not_mined_alongside_its_code_block() -> None:
    """A ``<div>`` wrapping ``<pre>`` used to look like a leaf.

    ``is_leaf_block`` searched a hand-written descendant list that omitted
    ``pre``/headings/cells, so the same code text was mined twice: once as
    NARRATIVE (sent to the model and injected into the finished book) and once
    as CODE (verbatim). ``<p>`` nesting already excluded the parent, so the fix
    is one vocabulary for both questions, not a new rule.
    """
    leaves = _html_leaves("<div><pre>int x = 1;</pre></div>")
    assert leaves == ["pre:int x = 1;"], leaves

    # Headings and definition lists behave the same way (both are BLOCK_TAGS).
    assert _html_leaves("<div><h2>Results</h2></div>") == ["h2:Results"]
    assert _html_leaves("<div><dl><dt>Term</dt><dd>Body</dd></dl></div>") == ["dt:Term", "dd:Body"]


def test_nested_check_uses_the_same_vocabulary_as_the_mining_set() -> None:
    """The leaf predicate and the mining set may never disagree again."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(
        "<div><pre>x</pre><h1>t</h1><table><tr><td>c</td></tr></table></div>", "html.parser"
    )
    block_names = set(BLOCK_TAGS)
    wrapper = soup.find("div")
    assert wrapper is not None
    assert not is_leaf_block(wrapper, block_names)


# --- Publication render must not mutate the caller's block list -------------


def test_diagram_vectorization_does_not_append_to_the_callers_list(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Export hands the same ``final_blocks`` to every render of the run.

    The academic-figure path appended ``pdf_main#fig_*`` straight into the
    caller's list, so a second render (visual-gate reflow, ``--emit-both``)
    shipped each figure twice under a duplicated block id.
    """
    import ubt.adapters.pdf.asset_extractor as asset_extractor
    import ubt.adapters.pdf.svg_diagram as svg_diagram
    from ubt.adapters.pdf.docling_render import DoclingRenderStrategy
    from ubt.core.ir.models import BoundingBox

    source_pdf = Path(__file__).resolve().parents[3] / "docs" / "synthetic-mono.pdf"
    if not source_pdf.is_file():
        pytest.skip(f"{source_pdf.name} fixture missing")

    monkeypatch.setattr(svg_diagram, "is_svg_backend_available", lambda: False)
    monkeypatch.setattr(svg_diagram, "is_svg_rendering_supported", lambda: False)
    monkeypatch.setattr(svg_diagram, "detect_diagram_regions", lambda *a, **k: [])

    fig_png = tmp_path / "fig.png"
    fig_png.write_bytes(b"\x89PNG\r\n\x1a\n")
    figure = asset_extractor.ExtractedFigure(
        fig_id="1",
        caption_en="FIG. 1: demo",
        page=1,
        image_path=fig_png,
        relative_path="assets/fig.png",
        bbox=(0.0, 0.0, 100.0, 100.0),
    )
    monkeypatch.setattr(asset_extractor, "extract_pdf_figures", lambda *a, **k: {"1": figure})

    strategy = DoclingRenderStrategy(
        reconstructor=SimpleNamespace(),  # type: ignore[arg-type]
        alternator=SimpleNamespace(),  # type: ignore[arg-type]
        diagram_localizer=cast("Any", SimpleNamespace(get_page_height=lambda _src, _page: 800.0)),
    )
    narrative = IRBlock(
        id="pdf_main#b001",
        spine_index=1,
        block_type=BlockType.NARRATIVE,
        flow_id=FlowID.MAIN_STORY,
        source_text="Body text.",
        bbox=BoundingBox(page=1, x0=0.0, y0=0.0, x1=100.0, y1=100.0),
    )
    blocks = [narrative]

    first, _ = strategy._vectorize_diagrams_sync(source_pdf, blocks, tmp_path / "assets", "zh")
    second, _ = strategy._vectorize_diagrams_sync(source_pdf, blocks, tmp_path / "assets", "zh")

    def figures(result: list[IRBlock]) -> list[str]:
        return [b.id for b in result if b.id.startswith("pdf_main#fig_")]

    assert len(blocks) == 1, "the caller's list must survive the render untouched"
    assert figures(first), "the fixture must actually exercise the figure-append path"
    assert figures(first) == figures(second), "two renders must produce two identical results"


# --- Localized caption labels ----------------------------------------------


@pytest.mark.parametrize(
    ("target_lang", "expected"),
    [
        ("zh", "表 3.1"),
        ("zh-tw", "表 3.1"),
        ("ja", "表 3.1"),
        ("ko", "표 3.1"),
        ("en", "Table 3.1"),
        ("fr", "Tableau 3.1"),
        ("de", "Tabelle 3.1"),
        ("es", "Tabla 3.1"),
        ("ru", "Таблица 3.1"),
    ],
)
def test_table_label_follows_the_target_language(target_lang: str, expected: str) -> None:
    """``TABLE 3.1`` was rewritten with a hard-coded Chinese literal.

    A French or English target book therefore shipped a Chinese "表" in its
    table captions, while the figure label beside it was already localized from
    the per-language profile.
    """
    from ubt.adapters.pdf.typst_reconstructor import _polish_target_text

    assert _polish_target_text("TABLE 3.1 Devices.", target_lang).startswith(expected)


def test_every_language_profile_names_both_labels() -> None:
    """A table prefix that silently equals the English default is the same bug."""
    for code in ("zh", "zh-tw", "ja", "ko", "fr", "de", "es", "ru"):
        config = resolve_font_config(code)
        assert config.table_prefix not in ("Table", ""), code
        assert config.figure_prefix not in ("Fig.", ""), code


# --- QE: the near-verbatim echo --------------------------------------------


def test_near_echo_is_the_same_defect_class_as_an_exact_echo() -> None:
    """An almost-verbatim target is an untranslated paragraph, not "other".

    The near-echo reason matched none of the score classifier's keywords, so it
    landed on ``QE_SCORE_STRUCTURAL_OTHER`` (0.70) — within 0.05 of the default
    threshold — and carried no structural marker, so the quality gate released
    a never-translated paragraph as ``MTQE_PASSED``.
    """
    near_echo = _SOURCE_PARAGRAPH.replace("layer", "layers")
    decision = FastPassFilter(target_lang="zh").evaluate(
        _SOURCE_PARAGRAPH, near_echo, block_type=BlockType.NARRATIVE
    )
    assert not decision.passed
    assert decision.reason.startswith(NEAR_ECHO_MARKER)
    assert HeuristicQERunner.score_from_decision_reason(decision.reason) == QE_SCORE_FABRICATED
    assert has_structural_defect([decision.reason])
    assert any(marker in decision.reason for marker in CRITICAL_DEFECT_MARKERS)


def test_both_echo_phrasings_are_registered_in_every_defect_table() -> None:
    """The reason strings and the tables must not be able to drift apart."""
    from ubt.core.qe.fast_pass import FastPassFilter as _FP

    for marker in (ECHO_MARKER, NEAR_ECHO_MARKER):
        assert marker in STRUCTURAL_DEFECT_MARKERS, marker
        assert marker in CRITICAL_DEFECT_MARKERS, marker

    exact = _FP(target_lang="zh").evaluate(
        _SOURCE_PARAGRAPH, _SOURCE_PARAGRAPH, block_type=BlockType.NARRATIVE
    )
    assert exact.reason.startswith(ECHO_MARKER)


# --- Preset ownership -------------------------------------------------------


def test_preset_engine_knobs_apply_only_after_an_explicit_pick() -> None:
    """An unpicked preset must not rewrite the user's ``UBT_*`` environment.

    A preset-applying surface used to merge
    ``PRESETS[STANDARD].engine_overrides()`` unconditionally, and overrides beat
    the environment — so ``UBT_RENDER_ENGINE`` / ``UBT_MATH_BACKEND`` /
    ``UBT_PROMPT_STRATEGY`` set by the operator were silently replaced on every
    run, the same class of bug the 2026-09 review removed for translate_chrome /
    cover_mode / formula_mode.

    The TUI surface this was originally written against no longer exists, so the
    invariant is asserted on the function that actually owns the decision —
    :func:`ubt.core.presets.resolve_engine_params`, the single place that ranks
    explicit flags over the preset bundle over the engine default. The
    assertions below are the original ones, unchanged in strength.
    """
    engine_keys = set(PRESETS[Preset.STANDARD].engine_overrides())
    assert engine_keys

    # Nothing picked, nothing passed: the resolver must inject nothing, so
    # ``UBT_*`` keeps precedence. (This is the regression the old assertion
    # ``not (engine_keys & set(untouched))`` pinned.)
    assert resolve_engine_params(None, dict.fromkeys(engine_keys)) == {}

    # An explicit pick contributes the whole bundle...
    resolved = resolve_engine_params(Preset.PUBLICATION, dict.fromkeys(engine_keys))
    assert engine_keys <= set(resolved)
    assert resolved["prompt_strategy"] == PRESETS[Preset.PUBLICATION].prompt_strategy

    # ...and an explicit flag still beats the bundle.
    assert (
        resolve_engine_params(Preset.PUBLICATION, {"prompt_strategy": "minimal"})["prompt_strategy"]
        == "minimal"
    )


# --- Export coverage gate ---------------------------------------------------


def _ledger_with_targets(tmp_path: Path, *, translated: int, total: int) -> SQLiteJobLedger:
    job_id = "job_coverage"
    ledger = SQLiteJobLedger(tmp_path / f"{job_id}.sqlite")
    blocks = [
        IRBlock(
            id=f"b{idx}",
            spine_index=idx,
            block_type=BlockType.NARRATIVE,
            flow_id=FlowID.MAIN_STORY,
            source_text=f"Source paragraph number {idx}.",
        )
        for idx in range(1, total + 1)
    ]
    ledger.init_job(
        job_id,
        DocumentIR(
            doc_id=job_id,
            source_path=str(tmp_path / "book.md"),
            format_type="markdown",
            blocks=blocks,
        ),
        target_lang="zh",
    )
    for block in blocks[:translated]:
        ledger.save_checkpoint(
            block_id=block.id,
            status=BlockStatus.MTQE_PASSED,
            target_text=f"译文 {block.id}",
        )
    for block in blocks[translated:]:
        ledger.save_checkpoint(block_id=block.id, status=BlockStatus.FAILED)
    return ledger


async def _run_export(ledger: SQLiteJobLedger, tmp_path: Path, *, ratio: float) -> Path:
    job_id = "job_coverage"
    manifest = BookManifest(
        doc_id=job_id,
        title="Coverage",
        source_path=str(tmp_path / "book.md"),
        target_lang="zh",
        source_lang="en",
    )
    (tmp_path / "book.md").write_text(
        "# Coverage\n\nSource paragraph number 1.\n", encoding="utf-8"
    )

    async def _event(*args: Any, **kwargs: Any) -> None:
        return None

    ctx = build_stage_ctx(
        tmp_path,
        # The coverage floor is a config knob now; the stage has no shadow
        # default of its own to keep in step with it.
        config=UBTConfig(db_dir=tmp_path, export_min_completion_ratio=ratio),
        ledger=ledger,
        job_id=job_id,
        manifest=manifest,
        adapter=MarkdownAdapter(),
        output_path=tmp_path / "out.md",
        input_path=tmp_path / "book.md",
        target_lang="zh",
        source_lang="en",
        glossary_dicts=[],
        html_validator=HTMLDeltaValidator(),
        create_event=_event,
    )
    _events = [e async for e in run_export_stage(ctx)]
    assert ctx.output_path is not None
    return Path(ctx.output_path)


@pytest.mark.asyncio
async def test_export_refuses_a_book_where_most_blocks_have_no_target(tmp_path: Path) -> None:
    """All-failed used to render source text and finalize the job as completed."""
    ledger = _ledger_with_targets(tmp_path, translated=1, total=4)
    with pytest.raises(IntegrityViolationError, match="carry a translation"):
        await _run_export(ledger, tmp_path, ratio=0.5)
    assert ledger.get_job_stats("job_coverage")["completed"] == 1


@pytest.mark.asyncio
async def test_coverage_gate_floor_is_configurable(tmp_path: Path) -> None:
    """The gate must be liftable to 0 for a knowingly-partial delivery."""
    ledger = _ledger_with_targets(tmp_path, translated=1, total=4)
    out = await _run_export(ledger, tmp_path, ratio=0.0)
    assert out.is_file()


# --- DOCX monolingual export -------------------------------------------------


def _build_docx_with_link_and_picture(path: Path) -> Path:
    """One linked-heading paragraph plus one paragraph carrying a picture."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    png = path.parent / "dot.png"
    png.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
            "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
        )
    )
    doc = Document()
    paragraph = doc.add_paragraph("original linked title")
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "_Toc1")
    for child in list(paragraph._p):
        if child.tag == qn("w:r"):
            paragraph._p.remove(child)
            link.append(child)
    paragraph._p.append(link)
    picture_para = doc.add_paragraph("see figure here")
    picture_para.add_run().add_picture(str(png))
    doc.save(str(path))
    return path


def _docx_body_facts(path: Path) -> tuple[list[str], int, int]:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    body = doc.element.body
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    return (
        texts,
        len(list(body.iter(qn("w:drawing")))),
        len(list(body.iter(qn("w:hyperlink")))),
    )


@pytest.mark.asyncio
async def test_docx_monolingual_keeps_pictures_and_links_and_drops_source(
    tmp_path: Path,
) -> None:
    """A "monolingual" DOCX still shipped the source sentence, minus its art.

    ``_replace_paragraph_in_place`` removed only ``w:r`` children: a hyperlink
    is a direct child of ``w:p``, so its source text survived the rewrite, while
    the runs it did remove were the ones carrying ``w:drawing`` — every inline
    picture in the paragraph went with them.
    """

    from ubt.adapters.docx.adapter import DOCXAdapter

    source = _build_docx_with_link_and_picture(tmp_path / "book.docx")
    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(source)
    blocks: list[IRBlock] = []
    async for chapter in adapter.parse_stream(source):
        blocks.extend(chapter.blocks)
    assert blocks, "the fixture must yield blocks"
    translated = [
        block.model_copy(update={"target_text": "译文本", "skip_translate": False})
        for block in blocks
    ]

    out = tmp_path / "mono.docx"
    await adapter.render_blocks(manifest, translated, "zh", out, bilingual_mode="monolingual")
    texts, drawings, hyperlinks = _docx_body_facts(out)
    assert all(text == "译文本" for text in texts), texts
    assert drawings == 1, "inline pictures must survive a monolingual rewrite"
    assert hyperlinks == 1, "the hyperlink must survive, wrapping the translation"

    both = tmp_path / "bilingual.docx"
    await adapter.render_blocks(manifest, translated, "zh", both, bilingual_mode="bilingual")
    _texts, bilingual_drawings, _links = _docx_body_facts(both)
    assert bilingual_drawings == 1


# --- Cost honesty in the delivery report ------------------------------------


def test_unpriced_model_reports_unknown_not_zero(tmp_path: Path) -> None:
    """A model missing from the price table must not read as a $0 delivery.

    The price table deliberately covers only the benchmarked models, so most
    operators' models are unpriced (the shipping default is explicitly priced at
    0.0). The pricing layer already returns ``None``; the report layer was
    the one collapsing it to 0.0, in a footnote that admitted "0 = unknown".
    """
    from ubt.core.engine.reporter import build_quality_report, render_kdp_audit_markdown

    ledger = _ledger_with_targets(tmp_path, translated=2, total=2)
    manifest = BookManifest(
        doc_id="job_coverage",
        title="Cost",
        source_path=str(tmp_path / "book.md"),
        target_lang="zh",
        source_lang="en",
    )

    unpriced = build_quality_report(ledger, "job_coverage", manifest, tmp_path / "out.md")
    assert unpriced.summary.estimated_cost_usd is None
    assert "unknown" in render_kdp_audit_markdown(unpriced)

    priced = build_quality_report(
        ledger, "job_coverage", manifest, tmp_path / "out.md", token_cost_usd=0.5
    )
    assert priced.summary.estimated_cost_usd == 0.5
    assert "$0.50000" in render_kdp_audit_markdown(priced)


def test_caption_driven_figure_crops_do_not_ship_the_same_figure_twice(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A caption rect that contains another crop printed one figure twice.

    On chapter-3-zh.pdf page 4, ``fig_3_3_p4`` fully contained ``fig_3_2_p4``
    (overlap/min-area 1.00) and overlapped ``fig_3_4_p4`` by 57%, so the reader
    got figure 3.2 again — with figure 3.2's caption inside figure 3.3 — while
    both were also emitted separately.
    """
    import ubt.adapters.pdf.asset_extractor as asset_extractor
    import ubt.adapters.pdf.svg_diagram as svg_diagram
    from ubt.adapters.pdf.docling_render import DoclingRenderStrategy
    from ubt.core.ir.models import BoundingBox

    source_pdf = Path(__file__).resolve().parents[3] / "docs" / "synthetic-mono.pdf"
    if not source_pdf.is_file():
        pytest.skip(f"{source_pdf.name} fixture missing")

    monkeypatch.setattr(svg_diagram, "is_svg_backend_available", lambda: False)
    monkeypatch.setattr(svg_diagram, "is_svg_rendering_supported", lambda: False)
    monkeypatch.setattr(svg_diagram, "detect_diagram_regions", lambda *a, **k: [])

    fig_png = tmp_path / "fig.png"
    fig_png.write_bytes(b"\x89PNG\r\n\x1a\n")

    def figure(fig_id: str, bbox: tuple[float, float, float, float]) -> object:
        return asset_extractor.ExtractedFigure(
            fig_id=fig_id,
            caption_en=f"FIG. {fig_id}: demo",
            page=1,
            image_path=fig_png,
            relative_path="assets/fig.png",
            bbox=bbox,
        )

    monkeypatch.setattr(
        asset_extractor,
        "extract_pdf_figures",
        lambda *a, **k: {
            "big": figure("big", (0.0, 0.0, 300.0, 300.0)),
            "inside": figure("inside", (100.0, 100.0, 200.0, 200.0)),
            "apart": figure("apart", (500.0, 500.0, 600.0, 600.0)),
        },
    )

    strategy = DoclingRenderStrategy(
        reconstructor=SimpleNamespace(),  # type: ignore[arg-type]
        alternator=SimpleNamespace(),  # type: ignore[arg-type]
        diagram_localizer=cast("Any", SimpleNamespace(get_page_height=lambda _src, _page: 800.0)),
    )
    blocks = [
        IRBlock(
            id="pdf_main#b001",
            spine_index=1,
            block_type=BlockType.NARRATIVE,
            flow_id=FlowID.MAIN_STORY,
            source_text="Body text.",
            bbox=BoundingBox(page=1, x0=0.0, y0=0.0, x1=100.0, y1=100.0),
        )
    ]

    out, _ = strategy._vectorize_diagrams_sync(source_pdf, blocks, tmp_path / "assets", "zh")
    ids = sorted(b.id for b in out if b.id.startswith("pdf_main#fig_"))
    assert ids == ["pdf_main#fig_apart", "pdf_main#fig_big"], ids
