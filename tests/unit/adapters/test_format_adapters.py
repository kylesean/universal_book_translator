"""The HTML/EPUB/DOCX adapters parse their format and render it back.

Parse: the adapter's ``parse_stream`` yields the expected block count for a
fixture whose every construct is known. Render: with every translatable block
given a ``tr:``-prefixed target, the monolingual render ships the target text
and the bilingual render pairs it with the source -- for DOCX including the
table, whose structured target is Markdown the adapter must lay back into real
table cells.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from ubt.adapters.docx.adapter import DOCXAdapter
from ubt.adapters.epub.adapter import EPUBAdapter
from ubt.adapters.html.adapter import HTMLAdapter
from ubt.core.ir.models import IRBlock

#: A 1x1 transparent PNG, for the picture fixtures.
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360606060000000050001a5f645400000000049454e44ae426082"
)

_HTML = """<!doctype html><html><head><title>t</title></head><body>
<h1>Chapter One</h1>
<p>Hello <b>world</b>, this is prose.</p>
<ul><li>First item</li><li>Second item</li></ul>
<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>
<figure><img src="pic.png"><figcaption>Figure 1</figcaption></figure>
<pre>print("hi")</pre>
</body></html>"""

_EPUB_CONTAINER = """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf"
    media-type="application/oebps-package+xml"/></rootfiles>
</container>"""

_EPUB_OPF = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>T</dc:title></metadata>
  <manifest>
    <item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>
    <item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>"""

_EPUB_CH1 = """<html><body><h1>One</h1><p>First chapter.</p>
<img src="images/fig.png" alt="fig"/></body></html>"""
_EPUB_CH2 = """<html><body><h1>Two</h1><p>Second chapter.</p>
<ul><li>a</li><li>b</li></ul></body></html>"""


def _write_epub(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr("META-INF/container.xml", _EPUB_CONTAINER)
        zf.writestr("OEBPS/content.opf", _EPUB_OPF)
        zf.writestr("OEBPS/ch1.xhtml", _EPUB_CH1)
        zf.writestr("OEBPS/ch2.xhtml", _EPUB_CH2)
        zf.writestr("OEBPS/images/fig.png", _PNG)


def _write_docx(path: Path) -> None:
    from docx import Document as DocxDocument

    image = path.with_suffix(".png")
    image.write_bytes(_PNG)
    doc = DocxDocument()
    doc.add_heading("Chapter One", level=1)
    doc.add_paragraph("Hello world, this is prose.")
    doc.add_paragraph("First item", style="List Bullet")
    doc.add_paragraph("Second item", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.add_picture(str(image))
    doc.save(str(path))


@pytest.fixture(scope="module")
def html_book(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("html") / "sample.html"
    path.write_text(_HTML, encoding="utf-8")
    return path


@pytest.fixture(scope="module")
def epub_book(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("epub") / "sample.epub"
    _write_epub(path)
    return path


@pytest.fixture(scope="module")
def docx_book(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("docx") / "sample.docx"
    _write_docx(path)
    return path


async def _collect(adapter: object, book: Path) -> list[IRBlock]:
    blocks: list[IRBlock] = []
    async for chapter in adapter.parse_stream(book):  # type: ignore[attr-defined]
        blocks.extend(chapter.blocks)
    return blocks


async def test_html_adapter_parses_and_renders(tmp_path: Path, html_book: Path) -> None:
    adapter = HTMLAdapter()
    manifest = await adapter.extract_manifest(html_book)
    blocks = await _collect(adapter, html_book)
    assert len(blocks) == 8, "the html fixture has exactly 8 blocks"
    for block in blocks:
        if not block.skip_translate:
            block.target_text = f"tr:{block.source_text}"

    mono = tmp_path / "mono.html"
    await adapter.render_blocks(manifest, blocks, "zh", mono, bilingual_mode="monolingual")
    assert mono.exists()
    assert "tr:Chapter One" in mono.read_text(encoding="utf-8")

    bilingual = tmp_path / "bi.html"
    await adapter.render_blocks(manifest, blocks, "zh", bilingual, bilingual_mode="bilingual")
    assert bilingual.exists()
    bilingual_text = bilingual.read_text(encoding="utf-8")
    assert "ubt-bilingual-target" in bilingual_text
    assert "Hello" in bilingual_text  # the source stays beside the target


async def test_epub_adapter_parses_and_renders(tmp_path: Path, epub_book: Path) -> None:
    adapter = EPUBAdapter()
    manifest = await adapter.extract_manifest(epub_book)
    blocks = await _collect(adapter, epub_book)
    assert len(blocks) == 6, "the epub fixture has exactly 6 blocks"
    for block in blocks:
        if not block.skip_translate:
            block.target_text = f"tr:{block.source_text}"

    for name, mode in (("mono.epub", "monolingual"), ("bi.epub", "bilingual")):
        out = tmp_path / name
        await adapter.render_blocks(manifest, blocks, "zh", out, bilingual_mode=mode)
        assert out.exists(), mode
        with zipfile.ZipFile(out) as zf:
            assert not zf.testzip(), f"{mode} output is not a valid zip"


async def test_docx_adapter_parses_and_renders(tmp_path: Path, docx_book: Path) -> None:
    from docx import Document as OpenDocx

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(docx_book)
    blocks = await _collect(adapter, docx_book)
    assert len(blocks) == 6, "the docx fixture has exactly 6 blocks"
    for block in blocks:
        if block.skip_translate:
            continue
        block.target_text = (
            "| tr:A | tr:B |\n| --- | --- |\n| tr:1 | tr:2 |"
            if block.block_type.value == "table"
            else f"tr:{block.source_text}"
        )

    mono = tmp_path / "mono.docx"
    await adapter.render_blocks(manifest, blocks, "zh", mono, bilingual_mode="monolingual")
    assert mono.exists()
    mono_doc = OpenDocx(str(mono))
    mono_text = " ".join(p.text for p in mono_doc.paragraphs)
    assert "tr:Chapter One" in mono_text
    mono_table = " ".join(cell.text for row in mono_doc.tables[0].rows for cell in row.cells)
    assert "tr:A" in mono_table

    bilingual = tmp_path / "bi.docx"
    await adapter.render_blocks(manifest, blocks, "zh", bilingual, bilingual_mode="bilingual")
    assert bilingual.exists()
    bi_doc = OpenDocx(str(bilingual))
    bi_text = " ".join(p.text for p in bi_doc.paragraphs)
    assert "Chapter One" in bi_text  # the source stays beside the target
    assert "tr:Chapter One" in bi_text


def test_epub_natural_sort_orders_embedded_numbers_numerically() -> None:
    from ubt.adapters.epub.adapter import _natural_sort_key

    names = ["ch10.xhtml", "ch2.xhtml", "ch1.xhtml"]
    assert sorted(names, key=_natural_sort_key) == ["ch1.xhtml", "ch2.xhtml", "ch10.xhtml"]
    # A plain lexicographic sort would put ch10 before ch2.
    assert sorted(names) != sorted(names, key=_natural_sort_key)
