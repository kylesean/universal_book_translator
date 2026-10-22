"""A DOCX whose body part declares a bomb-scale size is refused before parsing.

``python-docx`` opens the package with a plain ``zipfile`` and decompresses
whatever the central directory declares, so a crafted ``word/document.xml``
takes the process down inside ``Document(...)`` — the kernel OOM-kills it and
there is no traceback to diagnose from. ``ZipReadBudget`` already guards every
member UBT reads itself; this is the one path where the read happens inside a
third-party loader.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document

from ubt.adapters.docx.adapter import _guard_package_size
from ubt.core.exceptions import DocumentParseError
from ubt.core.zip_safety import MAX_MEMBER_BYTES

pytestmark = pytest.mark.fast


def _docx(tmp_path: Path, name: str = "base.docx") -> Path:
    doc = Document()
    doc.add_paragraph("hello")
    path = tmp_path / name
    doc.save(str(path))
    return path


def _with_body_of(tmp_path: Path, declared_bytes: int) -> Path:
    """A real DOCX whose ``word/document.xml`` declares ``declared_bytes``."""
    base = _docx(tmp_path)
    bomb = tmp_path / "bomb.docx"
    payload = (
        '<?xml version="1.0"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:p><w:r><w:t>{'A' * declared_bytes}</w:t></w:r></w:p></w:body>"
        "</w:document>"
    ).encode()
    with zipfile.ZipFile(base) as source, zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            data = payload if item.filename == "word/document.xml" else source.read(item.filename)
            out.writestr(item.filename, data)
    return bomb


def test_a_normal_docx_passes_the_guard(tmp_path: Path) -> None:
    _guard_package_size(_docx(tmp_path))


def test_a_bomb_shaped_body_is_refused(tmp_path: Path) -> None:
    bomb = _with_body_of(tmp_path, MAX_MEMBER_BYTES + 1024)
    # The bomb is tiny on disk: the danger is entirely in the declared size.
    assert bomb.stat().st_size < 2 * 1024 * 1024

    with pytest.raises(DocumentParseError, match="decompression cap"):
        _guard_package_size(bomb)


def test_the_refusal_names_the_member_and_the_size(tmp_path: Path) -> None:
    bomb = _with_body_of(tmp_path, MAX_MEMBER_BYTES + 1024)
    declared = zipfile.ZipFile(bomb).getinfo("word/document.xml").file_size

    with pytest.raises(DocumentParseError) as excinfo:
        _guard_package_size(bomb)

    message = str(excinfo.value)
    assert "word/document.xml" in message
    assert str(declared) in message
    assert str(MAX_MEMBER_BYTES) in message


def test_a_non_zip_file_is_left_to_python_docx(tmp_path: Path) -> None:
    """A corrupt file must keep python-docx's own message, not a zip error."""
    junk = tmp_path / "junk.docx"
    junk.write_bytes(b"not a zip archive")

    _guard_package_size(junk)  # no raise: the loader reports it
