"""Unit tests for the DOCX adapter."""

import base64
import tempfile
from collections.abc import AsyncIterator
from pathlib import Path

import pytest
from docx import Document

from ubt.adapters.docx.adapter import DOCXAdapter
from ubt.core.engine.ledger import SQLiteJobLedger
from ubt.core.ir.models import BlockStatus, BlockType, BookManifest, ChapterMeta, FlowID, IRBlock


async def _collect[T](agen: AsyncIterator[T]) -> list[T]:
    return [item async for item in agen]


def _make_doc(path: Path) -> None:
    doc = Document()
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("Revenue grew steadily this quarter.")
    doc.add_paragraph("")  # blank paragraph must be skipped
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "EPS"
    table.cell(1, 1).text = "1.23 USD"
    doc.save(str(path))


@pytest.mark.asyncio
async def test_docx_manifest_and_parse() -> None:
    tmp = Path(tempfile.mkdtemp())
    docx_path = tmp / "report.docx"
    _make_doc(docx_path)

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(docx_path)
    assert manifest.chapters[0].chapter_id == "docx_main"
    assert manifest.chapters[0].spine_index == 1

    chapters = await _collect(adapter.parse_stream(docx_path))
    assert len(chapters) == 1
    blocks = chapters[0].blocks
    # heading + 1 para + 4 table cells (blank paragraph skipped)
    assert len(blocks) == 6
    assert blocks[0].block_type == BlockType.HEADING
    assert blocks[1].block_type == BlockType.NARRATIVE
    assert all(b.flow_id == FlowID.TABLE_GRID for b in blocks[2:])
    # cell ids encode table/row/col/para positions
    assert blocks[2].id == "docx_main#t000r000c000p000"
    assert blocks[5].id == "docx_main#t000r001c001p000"
    assert blocks[5].source_text == "1.23 USD"


@pytest.mark.asyncio
async def test_docx_bilingual_render_preserves_source_document() -> None:
    tmp = Path(tempfile.mkdtemp())
    docx_path = tmp / "report.docx"
    _make_doc(docx_path)
    out_path = tmp / "report_bilingual.docx"

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(docx_path)
    chapters = await _collect(adapter.parse_stream(docx_path))

    with SQLiteJobLedger(tmp / "ledger.db") as ledger:
        ledger.init_job_from_manifest(manifest.doc_id, manifest)
        ledger.append_chapter(manifest.doc_id, chapters[0])
        for block in chapters[0].blocks:
            ledger.save_checkpoint(
                block_id=block.id,
                status=BlockStatus.MTQE_PASSED,
                target_text=f"译：{block.source_text}",
            )
        rendered = await adapter.render_output(
            manifest, ledger, target_lang="zh", output_path=out_path
        )
        assert rendered == out_path
        assert out_path.exists()

    # bilingual output: source paragraph followed by target paragraph
    out_doc = Document(str(out_path))
    para_texts = [p.text for p in out_doc.paragraphs]
    assert para_texts[0] == "Quarterly Report"
    assert "译：Quarterly Report" in para_texts[1]
    assert "Revenue grew steadily this quarter." in para_texts
    assert any("译：Revenue grew steadily this quarter." in t for t in para_texts)

    # table content preserved with translations inside cells
    assert len(out_doc.tables) == 1
    cell_texts = [c.text for r in out_doc.tables[0].rows for c in r.cells]
    assert any("译：Metric" in t for t in cell_texts)
    assert any("1.23 USD" in t for t in cell_texts)  # source cell intact

    # THE SOURCE DOCUMENT IS NEVER MODIFIED
    src_doc = Document(str(docx_path))
    assert [p.text for p in src_doc.paragraphs] == [
        "Quarterly Report",
        "Revenue grew steadily this quarter.",
        "",
    ]
    assert len(src_doc.tables) == 1


@pytest.mark.asyncio
async def test_docx_merged_cells_are_not_duplicated() -> None:
    """Merged cells must be mined exactly once (id() collision regression)."""
    tmp = Path(tempfile.mkdtemp())
    doc = Document()
    doc.add_paragraph("Intro")
    table = doc.add_table(rows=1, cols=3)
    # merge cells 1 and 2 of the single row
    a = table.cell(0, 1)
    b = table.cell(0, 2)
    merged = a.merge(b)
    merged.text = "Spanning"
    path = tmp / "merged.docx"
    doc.save(str(path))

    adapter = DOCXAdapter()
    chapters = await _collect(adapter.parse_stream(path))
    cell_blocks = [b for b in chapters[0].blocks if b.id.startswith("docx_main#t")]
    texts = [b.source_text for b in cell_blocks]
    assert texts == ["Spanning"]  # exactly once, not duplicated


@pytest.mark.asyncio
async def test_docx_blank_paragraphs_do_not_shift_block_ids() -> None:
    """Empty paragraphs must not consume a para_idx.

    Extraction (:meth:`DOCXAdapter._extract_blocks_sync`) skips empty
    paragraphs *without* incrementing ``para_idx``. If the render path
    increments anyway, every block after the first blank shifts by one and
    translations are silently injected into the wrong paragraph (or dropped
    entirely). The blank must sit in the middle — a trailing blank hides the
    bug, which is why the original fixture never caught it.
    """
    tmp = Path(tempfile.mkdtemp())
    doc = Document()
    doc.add_paragraph("First")
    doc.add_paragraph("")  # blank — must not consume an index
    doc.add_paragraph("")  # second blank — shifts by two if the bug returns
    doc.add_paragraph("Second")
    path = tmp / "blanks.docx"
    doc.save(str(path))

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(path)
    chapters = await _collect(adapter.parse_stream(path))
    blocks = chapters[0].blocks
    assert [b.source_text for b in blocks] == ["First", "Second"]
    assert [b.id for b in blocks] == ["docx_main#p00000", "docx_main#p00001"]

    for block in blocks:
        block.target_text = f"译：{block.source_text}"

    out = tmp / "blanks_bilingual.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out)

    texts = [p.text for p in Document(str(out)).paragraphs]
    assert texts == ["First", "译：First", "", "", "Second", "译：Second"]


@pytest.mark.asyncio
async def test_docx_table_cell_paragraph_index_mirror() -> None:
    """Regression H3: extraction and render must number cell paragraphs identically.

    Both sides ``enumerate(cell.paragraphs)`` and skip empties *after* computing
    the block id, so an empty leading paragraph consumes an index on both sides.
    If either side diverged (e.g. render skipped the empty without advancing),
    the translation would land in the wrong cell paragraph and silently corrupt
    the table. This locks the cell path (the main-paragraph path has its own
    blank-shift guard test above).
    """
    tmp = Path(tempfile.mkdtemp())
    doc = Document()
    doc.add_paragraph("Body")
    table = doc.add_table(rows=1, cols=2)
    c0 = table.cell(0, 0)
    c0.paragraphs[0].add_run("")  # leading empty paragraph (p0000)
    c0.add_paragraph("Hello")  # real content (p0001)
    table.cell(0, 1).paragraphs[0].add_run("World")  # p0000 of col 1
    path = tmp / "cells.docx"
    doc.save(str(path))

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(path)
    blocks = (await _collect(adapter.parse_stream(path)))[0].blocks
    cell_blocks = {b.id: b for b in blocks if b.id.startswith("docx_main#t")}

    # The empty leading paragraph must NOT take index 1; "Hello" is p001.
    assert cell_blocks["docx_main#t000r000c000p001"].source_text == "Hello"
    assert cell_blocks["docx_main#t000r000c001p000"].source_text == "World"

    for block in blocks:
        block.target_text = f"ZH:{block.source_text}"
    out = tmp / "cells_mono.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out, bilingual_mode="monolingual")

    out_cell = Document(str(out)).tables[0].cell(0, 0)
    # Empty stays empty; the translation landed on the correct (second) paragraph.
    assert [p.text for p in out_cell.paragraphs] == ["", "ZH:Hello"]
    assert Document(str(out)).tables[0].cell(0, 1).text == "ZH:World"


@pytest.mark.asyncio
async def test_docx_missing_file_raises() -> None:
    from ubt.core.exceptions import DocumentParseError

    adapter = DOCXAdapter()
    with pytest.raises(DocumentParseError, match="DOCX file not found"):
        await adapter.extract_manifest(Path("/nonexistent/file.docx"))


@pytest.mark.asyncio
async def test_docx_bilingual_render_parses_html_formatting_without_tag_leak() -> None:
    tmp = Path(tempfile.mkdtemp())
    doc = Document()
    doc.add_paragraph("Original text")
    path = tmp / "tags.docx"
    doc.save(str(path))

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(path)
    chapters = await _collect(adapter.parse_stream(path))
    blocks = chapters[0].blocks

    blocks[
        0
    ].target_text = '<mark class="ubt-warning">这是 <em>重要</em> 的 <strong>测试</strong></mark>'
    out = tmp / "tags_bilingual.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out)

    out_doc = Document(str(out))
    texts = [p.text for p in out_doc.paragraphs]
    # Raw HTML tags must not leak into visible paragraph text
    assert "<mark" not in texts[1]
    assert "<em>" not in texts[1]
    assert "<strong>" not in texts[1]
    assert texts[1] == "这是 重要 的 测试"
    # The formatted runs must have bold/italic attributes applied
    runs = out_doc.paragraphs[1].runs
    italic_runs = [r for r in runs if r.italic]
    bold_runs = [r for r in runs if r.bold]
    assert any("重要" in r.text for r in italic_runs)
    assert any("测试" in r.text for r in bold_runs)


@pytest.mark.asyncio
async def test_docx_render_preserves_angle_bracket_prose_without_content_loss() -> None:
    """Regression (H1): translated technical prose containing ``List<T>``,
    ``<stdio.h>`` or ``a<b`` must reach the Word paragraph intact instead of
    being swallowed as HTML tags by the run builder, while a genuine inline
    ``<b>`` still becomes a bold run."""
    tmp = Path(tempfile.mkdtemp())
    doc = Document()
    doc.add_paragraph("Original text")
    path = tmp / "angle.docx"
    doc.save(str(path))

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(path)
    chapters = await _collect(adapter.parse_stream(path))
    blocks = chapters[0].blocks
    blocks[0].target_text = "每个 List<T> 含 <stdio.h>，当 a<b 且 b>c，另有 <b>粗</b>"
    out = tmp / "angle_bilingual.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out)

    para = Document(str(out)).paragraphs[1]
    joined = "".join(r.text for r in para.runs)
    for piece in ("List<T>", "<stdio.h>", "a<b 且 b>c"):
        assert piece in joined, (piece, joined)
    assert any(r.bold and "粗" in r.text for r in para.runs), joined


@pytest.mark.asyncio
async def test_docx_nested_table_cells_are_extracted_and_translated(tmp_path: Path) -> None:
    """Regression (P1-b2): a table whose cell contains a *nested* table must
    have the nested cells' text enter the IR (previously only top-level tables
    and a cell's direct paragraphs were walked, so nested text silently shipped
    untranslated) and be translated on the way back out. The nested block id
    carries the ``n000`` nesting segment."""
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer_cell = outer.cell(0, 0)
    outer_cell.paragraphs[0].text = "OuterLabel"
    nested = outer_cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).paragraphs[0].text = "InnerValue"
    path = tmp_path / "nested.docx"
    doc.save(str(path))

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(path)
    chapters = await _collect(adapter.parse_stream(path))
    blocks = chapters[0].blocks
    sources = {b.source_text for b in blocks}
    assert "OuterLabel" in sources
    assert "InnerValue" in sources, "nested cell text never entered the IR"

    nested_block = next(b for b in blocks if b.source_text == "InnerValue")
    assert "#t000r000c000n000r000c000p000" in nested_block.id, nested_block.id

    # Round-trip: translate every block and confirm the nested cell got the
    # injected target — proves extraction and render share one id scheme.
    for b in blocks:
        b.target_text = f"ZH:{b.source_text}"
    out = tmp_path / "out.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out)

    out_doc = Document(str(out))
    nested_cell = out_doc.tables[0].cell(0, 0).tables[0].cell(0, 0)
    assert "ZH:InnerValue" in nested_cell.text, nested_cell.text


@pytest.mark.asyncio
async def test_docx_section_headers_and_footers_are_extracted_and_translated(
    tmp_path: Path,
) -> None:
    """Regression (P1-b3): section headers/footers carry translatable text that
    previously never reached the IR (the body walk only covered ``doc.element.body``).
    They must be extracted into blocks and translated back in place; inherited
    (linked-to-previous) parts must be skipped so a shared header isn't
    double-processed."""
    doc = Document()
    doc.add_paragraph("Body text here.")
    sec = doc.sections[0]
    sec.header.is_linked_to_previous = False
    sec.header.paragraphs[0].text = "Confidential Header"
    sec.footer.is_linked_to_previous = False
    sec.footer.paragraphs[0].text = "Page footer note"
    path = tmp_path / "hf.docx"
    doc.save(str(path))

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(path)
    chapters = await _collect(adapter.parse_stream(path))
    blocks = chapters[0].blocks
    sources = {b.source_text for b in blocks}
    assert "Confidential Header" in sources, sources
    assert "Page footer note" in sources, sources
    assert "Body text here." in sources

    # Header ids live in their own namespace and are distinct from body ids.
    hdr = next(b for b in blocks if b.source_text == "Confidential Header")
    assert "#h000hdr" in hdr.id, hdr.id

    for b in blocks:
        b.target_text = f"ZH:{b.source_text}"
    out = tmp_path / "hf_out.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out, bilingual_mode="monolingual")

    out_sec = Document(str(out)).sections[0]
    assert out_sec.header.paragraphs[0].text == "ZH:Confidential Header"
    assert out_sec.footer.paragraphs[0].text == "ZH:Page footer note"


@pytest.mark.asyncio
async def test_docx_footnotes_are_extracted_and_translated(tmp_path: Path) -> None:
    """Regression (P1-b4): footnote text lives in a separate ``word/footnotes.xml``
    part that python-docx exposes only as a raw blob — it previously never
    reached the IR, so footnotes silently shipped untranslated. They must be
    extracted into blocks and translated back into the part."""
    import io
    import re
    import zipfile

    from lxml import etree

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    RT_FN = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
    CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"

    doc = Document()
    para = doc.add_paragraph("Body text")
    run = para.add_run()
    ref = etree.SubElement(run._r, f"{{{W}}}footnoteReference")
    ref.set(f"{{{W}}}id", "2")
    buf = io.BytesIO()
    doc.save(buf)

    footnotes = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:footnotes xmlns:w="{W}">'
        f'<w:footnote w:id="-1"><w:p><w:r><w:footnoteRef/></w:r></w:p></w:footnote>'
        f'<w:footnote w:id="2"><w:p><w:r><w:footnoteRef/></w:r>'
        f"<w:r><w:t>Original footnote text.</w:t></w:r></w:p></w:footnote>"
        f"</w:footnotes>"
    ).encode()

    path = tmp_path / "notes.docx"
    with (
        zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin,
        zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    f'<Override PartName="/word/footnotes.xml" ContentType="{CT}"/></Types>'.encode(),
                )
            if it.filename == "word/_rels/document.xml.rels":
                data = re.sub(
                    rb"</Relationships>",
                    f'<Relationship Id="rIdFn" Type="{RT_FN}" Target="footnotes.xml"/></Relationships>'.encode(),
                    data,
                )
            zout.writestr(it, data)
        zout.writestr("word/footnotes.xml", footnotes)

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(path)
    chapters = await _collect(adapter.parse_stream(path))
    blocks = chapters[0].blocks
    sources = {b.source_text for b in blocks}
    assert "Body text" in sources
    assert "Original footnote text." in sources, sources  # the previously-dropped content

    note_block = next(b for b in blocks if b.source_text == "Original footnote text.")
    assert note_block.id == "docx_main#fn00002p00000", note_block.id

    for b in blocks:
        b.target_text = f"ZH:{b.source_text}"
    out = tmp_path / "notes_out.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out, bilingual_mode="monolingual")

    with zipfile.ZipFile(out) as z:
        fx = z.read("word/footnotes.xml").decode("utf-8")
    assert "ZH:Original footnote text." in fx, fx
    # The source <w:t> was blanked, so "Original footnote text." must only
    # remain as the tail of the prefixed translation, never as its own node.
    assert ">Original footnote text." not in fx, "monolingual should blank the source text"


def test_docx_declares_output_suffixes() -> None:
    """§10.3: an empty output_suffixes means "no validation", so a docx written
    to a .pdf/.txt path reported success on a mislabelled file. The pipeline
    guard (engine/pipeline.py getattr(type(adapter), "output_suffixes")) only
    fires when the adapter declares them."""
    assert DOCXAdapter.output_suffixes == frozenset({".docx"})


async def test_docx_adapter_monolingual_mode() -> None:
    """Verify that DOCXAdapter supports monolingual target mode by replacing text in-place."""
    tmp = Path(tempfile.mkdtemp())
    doc = Document()
    doc.add_paragraph("Original English Sentence.")
    doc_path = tmp / "src.docx"
    doc.save(str(doc_path))

    adapter = DOCXAdapter()
    manifest = BookManifest(
        doc_id="test_docx",
        title="Docx Test",
        source_path=str(doc_path),
        chapters=[
            ChapterMeta(
                chapter_id="docx_main", title="Docx Test", spine_index=1, source_file="src.docx"
            )
        ],
    )
    blocks = [
        IRBlock(
            id="docx_main#p00000",
            spine_index=1,
            flow_id=FlowID.MAIN_STORY,
            source_text="Original English Sentence.",
            target_text="原生中文翻译段落。",
        )
    ]

    out_mono = tmp / "mono.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out_mono, bilingual_mode="target")
    doc_result = Document(str(out_mono))
    paragraphs = [p.text for p in doc_result.paragraphs if p.text.strip()]
    assert paragraphs == ["原生中文翻译段落。"]


async def test_bilingual_render_does_not_duplicate_section_breaks(
    tmp_path: Path,
) -> None:
    """Word stores a section break inside the last paragraph's pPr.

    The translated clone copied that pPr verbatim, so a two-section book
    rendered bilingual delivered three sections — duplicated page size, margins
    and header/footer references in the middle of a chapter.
    """
    import zipfile

    import lxml.etree as ET
    from docx import Document
    from docx.oxml.ns import qn

    src = tmp_path / "book.docx"
    doc = Document()
    doc.add_paragraph("Chapter one opens here.")
    doc.add_section()
    doc.add_paragraph("Chapter two opens here.")
    for para in doc.paragraphs:
        ppr = para._p.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            para.add_run("End of chapter one.")
    doc.save(str(src))

    def breaks_in_paragraphs(path: Path) -> int:
        with zipfile.ZipFile(path) as z:
            root = ET.fromstring(z.read("word/document.xml"))
        return len(root.findall(".//" + qn("w:pPr") + "/" + qn("w:sectPr")))

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(src)
    chapters = await _collect(adapter.parse_stream(src))
    blocks = [b for ch in chapters for b in ch.blocks]
    assert any(b.source_text == "End of chapter one." for b in blocks)
    for b in blocks:
        b.target_text = "译：" + (b.source_text or "")

    out = tmp_path / "out.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out, bilingual_mode="bilingual")

    assert breaks_in_paragraphs(out) == breaks_in_paragraphs(src) == 1
    assert len(Document(str(out)).sections) == 2
    assert any(
        t.startswith("译：End of chapter one.")
        for t in (p.text for p in Document(str(out)).paragraphs)
    )


@pytest.mark.fast
def test_docx_footnote_translation_strips_mark_tags_and_pads() -> None:
    from docx.oxml.ns import qn
    from lxml import etree

    from ubt.adapters.docx.adapter import _apply_note_translation

    # Simulate docx footnote paragraph (using lxml Element)
    para_el = etree.Element(qn("w:p"))
    r = etree.SubElement(para_el, qn("w:r"))
    t1 = etree.SubElement(r, qn("w:t"))
    t1.text = "Original footnote."

    # Apply bilingual note translation containing <mark> tag
    translated = '<mark class="ubt-blocked-human">【待人工审校：注记】</mark>'
    _apply_note_translation(para_el, translated, is_monolingual=False)

    all_texts = [node.text for node in para_el.iter(qn("w:t")) if node.text]
    full_text = "".join(all_texts)

    # Should not contain raw mark tags
    assert "<mark" not in full_text
    # Should pad between original and translated note
    assert "Original footnote. " in full_text or "Original footnote.\n" in full_text
    assert "【待人工审校：注记】" in full_text


@pytest.mark.fast
def test_docx_populate_runs_sets_east_asia_font() -> None:
    """DOCX translated runs must set w:eastAsia font for CJK targets to ensure correct rendering."""
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph("Original English")
    p.runs[0].font.name = "Calibri"
    src_rpr = p.runs[0]._r.find(qn("w:rPr"))

    p2 = doc.add_paragraph()
    DOCXAdapter._populate_paragraph_runs(p2, "中文翻译内容", src_rpr, target_lang="zh")
    rpr = p2.runs[0]._r.find(qn("w:rPr"))
    assert rpr is not None
    rfonts = rpr.find(qn("w:rFonts"))
    assert rfonts is not None
    assert rfonts.get(qn("w:ascii")) == "Calibri"
    assert rfonts.get(qn("w:eastAsia")) == "SimSun"


def _a0920_build_docx_with_link_and_picture(path: Path) -> Path:
    """One linked-heading paragraph plus one paragraph carrying a picture."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    png = path.parent / "dot.png"
    png.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
            "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
        )
    )
    doc = Document()
    paragraph = doc.add_paragraph("original linked title")
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "_Toc1")
    for child in list(paragraph._p):
        if child.tag == qn("w:r"):
            paragraph._p.remove(child)
            link.append(child)
    paragraph._p.append(link)
    picture_para = doc.add_paragraph("see figure here")
    picture_para.add_run().add_picture(str(png))
    doc.save(str(path))
    return path


def _a0920_docx_body_facts(path: Path) -> tuple[list[str], int, int]:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    body = doc.element.body
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    return (
        texts,
        len(list(body.iter(qn("w:drawing")))),
        len(list(body.iter(qn("w:hyperlink")))),
    )


@pytest.mark.asyncio
async def test_docx_monolingual_keeps_pictures_and_links_and_drops_source(
    tmp_path: Path,
) -> None:
    """A "monolingual" DOCX still shipped the source sentence, minus its art.

    ``_replace_paragraph_in_place`` removed only ``w:r`` children: a hyperlink
    is a direct child of ``w:p``, so its source text survived the rewrite, while
    the runs it did remove were the ones carrying ``w:drawing`` — every inline
    picture in the paragraph went with them.
    """

    from ubt.adapters.docx.adapter import DOCXAdapter

    source = _a0920_build_docx_with_link_and_picture(tmp_path / "book.docx")
    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(source)
    blocks: list[IRBlock] = []
    async for chapter in adapter.parse_stream(source):
        blocks.extend(chapter.blocks)
    assert blocks, "the fixture must yield blocks"
    translated = [
        block.model_copy(update={"target_text": "译文本", "skip_translate": False})
        for block in blocks
    ]

    out = tmp_path / "mono.docx"
    await adapter.render_blocks(manifest, translated, "zh", out, bilingual_mode="monolingual")
    texts, drawings, hyperlinks = _a0920_docx_body_facts(out)
    assert all(text == "译文本" for text in texts), texts
    assert drawings == 1, "inline pictures must survive a monolingual rewrite"
    assert hyperlinks == 1, "the hyperlink must survive, wrapping the translation"

    both = tmp_path / "bilingual.docx"
    await adapter.render_blocks(manifest, translated, "zh", both, bilingual_mode="bilingual")
    _texts, bilingual_drawings, _links = _a0920_docx_body_facts(both)
    assert bilingual_drawings == 1


@pytest.mark.asyncio
async def test_docx_failed_block_stays_labelled(tmp_path: Path) -> None:
    """DOCX cannot store the export ``<mark>``; a failed draft must stay labelled.

    It used to strip the failure mark and ship the raw machine draft as a
    finished translation — indistinguishable from an approved one.
    """
    doc = Document()
    doc.add_paragraph("Original text")
    path = tmp_path / "failed.docx"
    doc.save(str(path))

    adapter = DOCXAdapter()
    manifest = await adapter.extract_manifest(path)
    blocks = (await _collect(adapter.parse_stream(path)))[0].blocks
    blocks[0].target_text = '<mark class="ubt-failed-draft" title="failed">機器草稿</mark>'
    blocks[0].status = BlockStatus.FAILED

    out = tmp_path / "failed_out.docx"
    await adapter.render_blocks(manifest, blocks, "zh", out, bilingual_mode="monolingual")

    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "[UBT]" in text, text
    assert "<mark" not in text, text


def test_docx_iter_body_items_walks_content_controls() -> None:
    """Paragraphs inside ``w:sdt/w:sdtContent`` must be extracted, not dropped.

    Content controls are nested under w:sdt rather than direct body children, so
    the old direct-iteration shipped them untranslated.
    """
    from docx.oxml.ns import qn

    from ubt.adapters.docx.adapter import _iter_body_items

    doc = Document()
    para = doc.add_paragraph("Inside a content control")
    body = doc.element.body
    sdt = body.makeelement(qn("w:sdt"), {})
    content = body.makeelement(qn("w:sdtContent"), {})
    body.remove(para._p)
    content.append(para._p)
    sdt.append(content)
    body.append(sdt)

    texts = [getattr(item, "text", "") for item in _iter_body_items(doc)]
    assert any("Inside a content control" in t for t in texts)
