"""The native HTML/EPUB/DOCX readers.

Each reader must give every element a real ``Span.chars`` that slices exactly
its own text out of ``CanonicalSource.text`` -- the reader's core promise, a
*verifiable* span -- and the ``Document`` must round-trip through the bridge
losslessly. The fixtures are built here (a minimal DOCX via python-docx, an
EPUB via zipfile, an HTML string), so the tests need no external sample files.
"""

from __future__ import annotations

import zipfile
from collections import Counter
from collections.abc import Callable
from pathlib import Path

import pytest

from ubt.analyze.assemble import element_text
from ubt.analyze.bridge import blocks_from_document, document_from_blocks
from ubt.analyze.reader_docx import read_docx
from ubt.analyze.reader_epub import read_epub
from ubt.analyze.reader_html import read_html
from ubt.model.ast import Document, Table

pytestmark = pytest.mark.fast

#: A 1x1 transparent PNG, for the DOCX picture fixture.
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360606060000000050001a5f645400000000049454e44ae426082"
)

_HTML = """<!doctype html><html><head><title>t</title></head><body>
<h1>Chapter One</h1>
<p>Hello <b>world</b>, this is prose.</p>
<ul><li>First item</li><li>Second item</li></ul>
<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>
<figure><img src="pic.png"><figcaption>Figure 1</figcaption></figure>
<pre>print("hi")</pre>
</body></html>"""

_EPUB_CONTAINER = """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf"
    media-type="application/oebps-package+xml"/></rootfiles>
</container>"""

_EPUB_OPF = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>T</dc:title></metadata>
  <manifest>
    <item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>
    <item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>"""

_EPUB_CH1 = """<html><body><h1>One</h1><p>First chapter.</p>
<img src="images/fig.png" alt="fig"/></body></html>"""
_EPUB_CH2 = """<html><body><h1>Two</h1><p>Second chapter.</p>
<ul><li>a</li><li>b</li></ul></body></html>"""


def _write_epub(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr("META-INF/container.xml", _EPUB_CONTAINER)
        zf.writestr("OEBPS/content.opf", _EPUB_OPF)
        zf.writestr("OEBPS/ch1.xhtml", _EPUB_CH1)
        zf.writestr("OEBPS/ch2.xhtml", _EPUB_CH2)
        zf.writestr("OEBPS/images/fig.png", _PNG)


def _write_docx(path: Path) -> None:
    from docx import Document as DocxDocument

    image = path.with_suffix(".png")
    image.write_bytes(_PNG)
    doc = DocxDocument()
    doc.add_heading("Chapter One", level=1)
    doc.add_paragraph("Hello world, this is prose.")
    doc.add_paragraph("First item", style="List Bullet")
    doc.add_paragraph("Second item", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.add_picture(str(image))
    doc.save(str(path))


def _build_html(tmp_path: Path) -> Path:
    path = tmp_path / "sample.html"
    path.write_text(_HTML, encoding="utf-8")
    return path


def _build_epub(tmp_path: Path) -> Path:
    path = tmp_path / "sample.epub"
    _write_epub(path)
    return path


def _build_docx(tmp_path: Path) -> Path:
    path = tmp_path / "sample.docx"
    _write_docx(path)
    return path


Case = tuple[str, Callable[[Path], Document], dict[str, int]]

_CASES: tuple[Case, ...] = (
    (
        "html",
        lambda path: read_html(path),
        {
            "heading": 1,
            "paragraph": 1,
            "list_item": 2,
            "table": 1,
            "figure": 1,
            "caption": 1,
            "code_block": 1,
        },
    ),
    (
        "epub",
        lambda path: read_epub(path),
        {"heading": 2, "paragraph": 2, "figure": 1, "list_item": 2},
    ),
    (
        "docx",
        lambda path: read_docx(path),
        {"heading": 1, "paragraph": 1, "list_item": 2, "table": 1, "figure": 1},
    ),
)
_BUILDERS: dict[str, Callable[[Path], Path]] = {
    "html": _build_html,
    "epub": _build_epub,
    "docx": _build_docx,
}


def _span_problems(document: Document) -> list[str]:
    problems: list[str] = []
    for element in document.elements:
        chars = element.span.chars
        if chars is None:
            problems.append(f"{element.id}: missing Span.chars")
            continue
        expected = element_text(element)
        got = document.source.text[chars[0] : chars[1]]
        if got != expected:
            problems.append(f"{element.id}: span slice {got[:30]!r} != {expected[:30]!r}")
    return problems


@pytest.mark.parametrize(("name", "read", "expect"), _CASES, ids=[c[0] for c in _CASES])
def test_every_element_has_an_exact_verifiable_span(
    tmp_path: Path, name: str, read: Callable[[Path], Document], expect: dict[str, int]
) -> None:
    document = read(_BUILDERS[name](tmp_path))
    assert document.elements, f"{name}: the reader extracted nothing"
    assert _span_problems(document) == []


@pytest.mark.parametrize(("name", "read", "expect"), _CASES, ids=[c[0] for c in _CASES])
def test_the_document_round_trips_through_the_bridge(
    tmp_path: Path, name: str, read: Callable[[Path], Document], expect: dict[str, int]
) -> None:
    document = read(_BUILDERS[name](tmp_path))
    round_tripped = document_from_blocks(
        blocks_from_document(document), doc_id=document.source.doc_id
    )
    assert len(round_tripped.elements) == len(document.elements)
    assert Counter(e.kind.value for e in round_tripped.elements) == Counter(
        e.kind.value for e in document.elements
    )


@pytest.mark.parametrize(("name", "read", "expect"), _CASES, ids=[c[0] for c in _CASES])
def test_the_expected_structure_is_present(
    tmp_path: Path, name: str, read: Callable[[Path], Document], expect: dict[str, int]
) -> None:
    document = read(_BUILDERS[name](tmp_path))
    counts = Counter(element.kind.value for element in document.elements)
    for kind, minimum in expect.items():
        assert counts.get(kind, 0) >= minimum, f"{name}: expected >= {minimum} {kind}"


def test_docx_vertically_merged_cell_is_emitted_once(tmp_path: Path) -> None:
    # python-docx resolves a vMerge continuation row's cell to the restart cell
    # above, so a per-row dedup repeated the merged value on every row. A
    # table-scoped set emits it once and blanks the continuations.
    from docx import Document as DocxDocument

    path = tmp_path / "vmerge.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Header"
    table.cell(0, 1).text = "r0"
    table.cell(1, 1).text = "r1"
    table.cell(2, 1).text = "r2"
    table.cell(0, 0).merge(table.cell(2, 0))
    doc.save(str(path))

    document = read_docx(path)
    tables = [e for e in document.elements if isinstance(e, Table)]
    assert len(tables) == 1
    markup = tables[0].markup
    assert markup.count("Header") == 1
    assert "| Header | r0 |" in markup
    assert "|  | r1 |" in markup


def test_a_markdown_blockquote_marker_never_reaches_the_text(tmp_path: Path) -> None:
    """`>` is Markdown syntax, not prose.

    Every other Markdown construct had a structural branch; a blockquote fell
    through to ``paragraph.append(line)``, so "> quoted" was sent to the model
    with its marker. The target then carried a stray angle bracket, and the
    delivered paragraph had a callout the renderers never understood. The
    element union has no quote type, so the faithful minimum is to translate the
    content and drop the marker — including inside a quote, where a heading or
    list must still be recognised.
    """
    from ubt.analyze.reader_md import read_md

    path = tmp_path / "quoted.md"
    path.write_text(
        "> 引用的句子\n> 第二行\n\n> # 引用里的标题\n\n> - 引用里的列表\n\n普通段落。\n",
        encoding="utf-8",
    )
    elements = read_md(path).elements

    texts = [getattr(el, "text", "") for el in elements]
    assert texts[0] == "引用的句子\n第二行"
    assert not any(text.strip().startswith(">") for text in texts), texts
    # Structure inside the quote survives. ``elements`` is a union of every
    # element type and only some carry ``text``, so the type is asserted first
    # and the attribute read through ``getattr`` (mypy --strict rejects a bare
    # ``elements[1].text`` on the union).
    assert type(elements[1]).__name__ == "Heading"
    assert getattr(elements[1], "text", None) == "引用里的标题"
    assert type(elements[2]).__name__ == "ListItem"
    assert getattr(elements[2], "text", None) == "引用里的列表"
    assert getattr(elements[3], "text", None) == "普通段落。"


def test_nested_and_spaced_blockquote_markers_are_stripped(tmp_path: Path) -> None:
    from ubt.analyze.reader_md import read_md

    path = tmp_path / "nested.md"
    path.write_text("> > deep quote\n>\t tabbed quote\n", encoding="utf-8")
    elements = read_md(path).elements
    text = getattr(elements[0], "text", "")
    assert "deep quote" in text and "tabbed quote" in text
    assert ">" not in text


def _minimal_epub(path: Path, *, with_nav: bool = True, repeated: bool = False) -> Path:
    """A minimal but valid EPUB 3 container, written by hand."""
    import zipfile

    nav_item = (
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
        if with_nav
        else ""
    )
    nav_ref = '<itemref idref="nav"/>' if with_nav else ""
    extra_ref = '<itemref idref="c1"/>' if repeated else ""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        archive.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>T</dc:title></metadata>'
            "<manifest>"
            '<item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>'
            '<item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>'
            f"{nav_item}"
            "</manifest>"
            f'<spine><itemref idref="c1"/><itemref idref="c2"/>{nav_ref}{extra_ref}</spine>'
            "</package>",
        )
        archive.writestr(
            "OEBPS/ch1.xhtml", "<html><body><p>Chapter one body text.</p></body></html>"
        )
        archive.writestr(
            "OEBPS/ch2.xhtml", "<html><body><p>Chapter two body text.</p></body></html>"
        )
        if with_nav:
            archive.writestr(
                "OEBPS/nav.xhtml",
                "<html><body><nav><ol><li>One</li><li>Two</li></ol></nav></body></html>",
            )
    return path


def test_the_native_reader_and_the_adapter_agree_on_the_spine() -> None:
    """One spine rule, not two.

    The adapter's copy filtered the navigation document and collapsed repeated
    spine members; the native reader's copy did neither. The same minimal EPUB
    therefore produced three pages through ``read_epub`` — the nav document
    counted as prose — and two chapters through the adapter. Both now call
    ``spine_chapters``.
    """
    import asyncio
    import tempfile
    import zipfile

    from ubt.adapters.epub.adapter import EPUBAdapter
    from ubt.analyze.reader_epub import read_epub, spine_chapters

    with tempfile.TemporaryDirectory() as tmp:
        path = _minimal_epub(Path(tmp) / "nav.epub", with_nav=True)
        # The reader sees exactly the two chapters.
        assert len(read_epub(path).elements) == 2

        # The adapter names exactly the same members, in the same order.
        with zipfile.ZipFile(path) as archive:
            shared = spine_chapters(archive, "OEBPS/content.opf")
        manifest = asyncio.run(EPUBAdapter().extract_manifest(path))
        assert [c.source_file for c in manifest.chapters] == shared
        assert shared == ["OEBPS/ch1.xhtml", "OEBPS/ch2.xhtml"]
        assert not any("nav" in member for member in shared), "the nav document is not prose"


def test_a_repeated_spine_member_contributes_one_chapter() -> None:
    """The same file twice would clone identical blocks and be paid twice."""
    import tempfile
    import zipfile

    from ubt.analyze.reader_epub import spine_chapters

    with tempfile.TemporaryDirectory() as tmp:
        path = _minimal_epub(Path(tmp) / "dup.epub", with_nav=False, repeated=True)
        with zipfile.ZipFile(path) as archive:
            members = spine_chapters(archive, "OEBPS/content.opf")
    assert members == ["OEBPS/ch1.xhtml", "OEBPS/ch2.xhtml"]
